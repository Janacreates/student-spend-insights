"""Build a short findings deck (PowerPoint) from reports/metrics.json and the saved figures."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from .config import Paths

NAVY = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x55, 0x5B, 0x63)


def _text(slide, left, top, width, height, lines, size=16, bold_first=False, color=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.bold = bool(bold_first and i == 0)
        p.font.color.rgb = color or GREY
        p.space_after = Pt(8)


def _slide(prs, title):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _text(s, 0.6, 0.35, 12.1, 1.0, [title], size=26, bold_first=True, color=NAVY)
    return s


def _pic(slide, path: Path, left, top, width):
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))


def build_deck(paths: Paths, m: dict) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    fig = paths.figures
    e, f, a, n, c = m["eda"], m["forecast"], m["anomalies"], m["nudges"], m["categorization"]

    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, 0.8, 2.4, 11.5, 1.2, ["Student spending: where it goes and how to steer it"], size=36, bold_first=True, color=NAVY)
    _text(s, 0.8, 3.7, 11.5, 1.5, [
        "Findings from a synthetic dataset of "
        f"{m['data']['students']} students over {m['data']['months']} months.",
        "Synthetic data: results describe the analysis pipeline, not real students."], size=18)

    s = _slide(prs, f"Half of spending is fixed; groceries and dining lead the rest")
    _pic(s, fig / "fig_spend_by_category.png", 0.6, 1.4, 7.0)
    top2 = " and ".join(e["top2_variable_categories"])
    _text(s, 8.0, 1.7, 4.9, 4.5, [
        f"Rent and tuition are {e['fixed_pct_of_all_spend']:.0f}% of all spend.",
        f"Of the controllable part, {top2} make up {e['top2_variable_share_pct']:.0f}%.",
        f"Median student: ${e['student_monthly_variable_spend']['median']:,.0f}/month variable spend "
        f"(p10 ${e['student_monthly_variable_spend']['p10']:,.0f}, p90 ${e['student_monthly_variable_spend']['p90']:,.0f}).",
    ])

    d = e["dining_exam_effect"]
    s = _slide(prs, f"Dining spend rises about {d['relative_lift_pct']:.0f}% in exam months")
    _pic(s, fig / "fig_exam_dining.png", 0.6, 1.4, 6.2)
    _text(s, 7.3, 1.7, 5.6, 4.5, [
        f"+${d['mean_diff']:.0f}/month per student (95% CI ${d['ci95_low']:.0f} to ${d['ci95_high']:.0f}).",
        f"Paired Wilcoxon signed-rank across {d['n_students']} students, p = {d['wilcoxon_p']:.1e}.",
        "Implication: budgets should flex by month rather than stay flat.",
        "The effect is built into the simulator, so this shows the method, not a real-world finding.",
    ])

    s = _slide(prs, "Forecasts beat simple baselines, with intervals that hold up")
    _pic(s, fig / "fig_forecast_accuracy.png", 0.4, 1.4, 6.2)
    _pic(s, fig / "fig_interval_coverage.png", 6.8, 1.4, 5.6)
    _text(s, 0.6, 5.5, 12.0, 1.6, [
        f"Gradient boosting MAE ${f['mae']['gradient_boosting']:.0f} vs ${min(v for k, v in f['mae'].items() if k != 'gradient_boosting'):.0f} "
        f"for the best baseline ({f['mae_improvement_vs_best_baseline_pct']:.0f}% better), walk-forward, no leakage.",
        f"Raw 80% intervals covered {100 * f['interval_coverage_raw']:.1f}%; after conformal calibration "
        f"{100 * f['interval_coverage_conformal']:.1f}%.",
    ], size=15)

    t = n["eval_at_tau"]
    sim = n["simulated_experiment"]
    s = _slide(prs, "Nudges: a usable trigger, and an honest test design")
    _pic(s, fig / "fig_nudge_precision_recall.png", 0.4, 1.4, 6.0)
    _pic(s, fig / "fig_ab_power.png", 6.7, 1.4, 6.0)
    _text(s, 0.6, 5.3, 12.0, 1.9, [
        f"Day-15 pace trigger (tau {n['tau_tuned_on_first_half']}): precision {100 * t['precision']:.0f}% and recall "
        f"{100 * t['recall']:.0f}% on held-out months, vs {100 * n['base_rate_overspend_eval']:.0f}% if everyone were nudged.",
        f"Simulated A/B (assumed 10% effect): needs about {sim['n_per_arm_needed_80pct_power_raw']:,} cases per arm, "
        f"or {sim['n_per_arm_needed_80pct_power_cuped']:,} with CUPED variance reduction.",
    ], size=15)

    s = _slide(prs, "Limitations and next steps")
    r = c["unseen_merchants"]
    _text(s, 0.8, 1.6, 11.7, 5.4, [
        "Data is synthetic; every effect size and rate reflects simulator assumptions.",
        f"Categorisation on merchants never seen in training: hybrid covers {100 * r['hybrid_coverage']:.0f}% "
        f"at {100 * r['hybrid_accuracy_on_covered']:.0f}% accuracy; the rest goes to review.",
        f"Anomaly rules catch {100 * a['rules']['recall']:.0f}% of injected anomalies but only "
        f"{100 * a['rules']['precision']:.0f}% of flags are real: review queues need triage.",
        "Next: pretrained text embeddings or an LLM for unseen merchants; validate on real, consented data.",
    ], size=18)

    out = paths.reports / "findings_deck.pptx"
    prs.save(out)
    return out
